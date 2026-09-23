import datetime
import openpyxl
import csv

TARGET_MATURITY = 10
SHEET_NAME = "4. nominal spot curve"
MATURITY_ROW = 4
DATE_COL = 1
DATA_START_ROW = 6


def extract_from_file(path, target_maturity=TARGET_MATURITY):
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Лист '{SHEET_NAME}' не найден в {path}. "
                          f"Доступные листы: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    target_col = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=MATURITY_ROW, column=c).value
        if isinstance(v, (int, float)) and abs(float(v) - target_maturity) < 1e-9:
            target_col = c
            break
    if target_col is None:
        raise ValueError(f"Не найден срок {target_maturity} лет в файле {path}")

    rows = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        d = ws.cell(row=r, column=DATE_COL).value
        val = ws.cell(row=r, column=target_col).value
        if isinstance(d, datetime.datetime) and val is not None:
            rows.append((d.date().isoformat(), val))
    return rows


def run(files, output_path, maturity=10.0):
    all_rows = {}
    for f in files:
        rows = extract_from_file(f, maturity)
        print(f"  извлечено {len(rows)} наблюдений")
        for date_str, val in rows:
            all_rows[date_str] = val  

    sorted_dates = sorted(all_rows.keys())
    with open(output_path, "w", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["date", f"yield_{maturity}y"])
        for d in sorted_dates:
            writer.writerow([d, all_rows[d]])

    print(f"\nГотово. Всего {len(sorted_dates)} наблюдений сохранено в {output_path}")
    print(f"Диапазон дат: {sorted_dates[0]} — {sorted_dates[-1]}")


if __name__ == "__main__":

    FILES = [
        "GLC Nominal daily data_1990 to 1994.xlsx",
        "GLC Nominal daily data_1995 to 1999.xlsx",
        "GLC Nominal daily data_2000 to 2004.xlsx",
    ]

    OUTPUT_PATH = "uk_10y_full.csv"

    MATURITY = 10.0  

    run(FILES, OUTPUT_PATH, MATURITY)